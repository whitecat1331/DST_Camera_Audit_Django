"""Tests for the rejection CSV report parsing, matching, and output builder."""

from __future__ import annotations

import tempfile
from pathlib import Path

from django.test import TestCase

from cameras.models import Installation
from services.rejection_csv import parse_rejection_csv
from services.rejection_match import (
    match_rejection_rows,
    resolve_ds_vnc_target,
    resolve_fx_tv_targets,
    serial_family,
)
from services.rejection_report import ReportRow, build_rejection_report

CSV = (
    "Row Label,Count of Reason,Rejection Reason\n"
    "DS012167,1,No FIM Video\n"
    "FX1263,1,No FIM Video\n"
    "ZZ999999,2,No FIM Video\n"
)


class ParseRejectionCsvTests(TestCase):
    def test_parses_header_and_rows(self):
        rows = parse_rejection_csv(CSV)
        self.assertEqual([r.serial for r in rows], ["DS012167", "FX1263", "ZZ999999"])
        self.assertEqual(rows[0].count, "1")
        self.assertEqual(rows[0].reason, "No FIM Video")
        self.assertEqual(rows[0].source_row, 2)

    def test_serial_family(self):
        self.assertEqual(serial_family("DS012167"), "ds")
        self.assertEqual(serial_family("fx1263"), "fx")
        self.assertEqual(serial_family("ABC"), "")


class MatchRejectionRowsTests(TestCase):
    def setUp(self):
        self.ds = Installation.objects.create(
            ims_id=1,
            identifier="I-1",
            primary_platform="LTI",
            pole_number="123456",
            camera_a="DS012167",
            camera_b="DS012168",
            tf_a_ip="10.2.3.14",
            tf_b_ip="10.2.3.15",
            is_active=True,
        )
        self.fx = Installation.objects.create(
            ims_id=2,
            identifier="I-DE-2",
            primary_platform="DE",
            pole_number="266135",
            serial_number="FX1263",
            is_active=True,
        )

    def test_matches_fx_and_ds_and_unmatched(self):
        report = match_rejection_rows(parse_rejection_csv(CSV))
        by_serial = {r.row.serial: r for r in report.rows}
        self.assertEqual(by_serial["DS012167"].status, "matched")
        self.assertEqual(by_serial["DS012167"].kind, "lti")
        self.assertEqual(by_serial["FX1263"].status, "matched")
        self.assertEqual(by_serial["FX1263"].kind, "de")
        self.assertEqual(by_serial["ZZ999999"].status, "unmatched")

    def test_dedupes_second_match(self):
        rows = parse_rejection_csv(
            "Row Label,Count of Reason,Rejection Reason\n"
            "DS012167,1,No FIM Video\n"
            "DS012167,2,No FIM Video\n"
        )
        report = match_rejection_rows(rows)
        statuses = [r.status for r in report.rows]
        self.assertEqual(statuses, ["matched", "duplicate"])

    def test_resolve_ds_vnc_target(self):
        host, lane, thumb = resolve_ds_vnc_target(self.ds, "DS012167")
        self.assertEqual(host, "10.2.3.14")
        self.assertEqual(lane, "vnc_l1")
        self.assertEqual(thumb, "vnc_l1")

    def test_resolve_fx_tv_targets_empty_without_mapping(self):
        # No DragonEyeTeamViewerId rows → no capture targets.
        self.assertEqual(resolve_fx_tv_targets(self.fx, "FX1263"), [])


class BuildRejectionReportTests(TestCase):
    def test_builds_xlsx_with_embedded_image(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            shot = tmp_path / "shot.png"
            Image.new("RGB", (320, 240), (255, 0, 0)).save(shot)
            rows = [
                ReportRow(
                    serial="DS012167",
                    count="1",
                    reason="No FIM Video",
                    result="OK",
                    screenshot_path=str(shot),
                )
            ]
            xlsx = tmp_path / "report.xlsx"
            built = build_rejection_report(rows, xlsx)
            self.assertTrue(xlsx.exists())
            self.assertEqual(built, xlsx)
            # The first three columns plus Result are written as cell values.
            from openpyxl import load_workbook

            wb = load_workbook(xlsx)
            ws = wb.active
            self.assertEqual(ws["A2"].value, "DS012167")
            self.assertEqual(ws["B2"].value, "1")
            self.assertEqual(ws["C2"].value, "No FIM Video")
            self.assertEqual(ws["D2"].value, "OK")
            self.assertTrue(ws._images)  # embedded screenshot present


class RejectionViewTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from django.test import Client

        from cameras.models import UserProfile

        self.user = User.objects.create_user("tech", password="pw")
        UserProfile.objects.create(user=self.user, ims_role="technician")
        self.client = Client()
        self.client.force_login(self.user)
        Installation.objects.create(
            ims_id=1,
            identifier="I-1",
            primary_platform="LTI",
            pole_number="123456",
            camera_a="DS012167",
            camera_b="DS012168",
            tf_a_ip="10.2.3.14",
            tf_b_ip="10.2.3.15",
            is_active=True,
        )
        Installation.objects.create(
            ims_id=2,
            identifier="I-DE-2",
            primary_platform="DE",
            pole_number="266135",
            serial_number="FX1263",
            is_active=True,
        )

    def test_page_renders(self):
        resp = self.client.get("/audits/rejection-report/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Rejection Report")

    def test_preview_matches(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            "rejections.csv",
            CSV.encode("utf-8"),
            content_type="text/csv",
        )
        resp = self.client.post(
            "/audits/rejection-report/preview/",
            {"file": upload},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["matched_count"], 2)
        self.assertEqual(data["unmatched_count"], 1)
